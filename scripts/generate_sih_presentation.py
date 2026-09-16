import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_sih_presentation(output_path, diagram_img_path):
    prs = Presentation()
    # 16:9 Widescreen layout (13.333" x 7.5")
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_DARK = RGBColor(15, 23, 42)
    COLOR_PRIMARY = RGBColor(30, 58, 138)
    COLOR_TEXT = RGBColor(30, 41, 59)
    COLOR_MUTED = RGBColor(71, 85, 105)

    slide = prs.slides.add_slide(blank_layout)

    # Background White
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_WHITE

    # -------------------------------------------------------------------------
    # TOP HEADER BANNER
    # -------------------------------------------------------------------------

    # Top Left Team Oval Badge
    team_oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.4), Inches(0.2), Inches(1.5), Inches(0.85))
    team_oval.fill.solid()
    team_oval.fill.fore_color.rgb = RGBColor(248, 250, 252)
    team_oval.line.color.rgb = RGBColor(147, 51, 234)
    team_oval.line.width = Pt(2)
    tf_team = team_oval.text_frame
    tf_team.word_wrap = True
    p_team = tf_team.paragraphs[0]
    p_team.text = "Your Team\nName"
    p_team.font.size = Pt(12)
    p_team.font.bold = True
    p_team.font.color.rgb = RGBColor(147, 51, 234)

    # Center Title Header
    title_box = slide.shapes.add_textbox(Inches(2.1), Inches(0.25), Inches(8.5), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_t = tf_title.paragraphs[0]
    p_t.text = "MAARAKSHA - Every Mother’s Health Matters"
    p_t.font.size = Pt(26)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_DARK
    p_t.font.name = "Georgia"

    # Top Right SIH 2026 Badge
    sih_box = slide.shapes.add_textbox(Inches(10.8), Inches(0.15), Inches(2.2), Inches(0.95))
    tf_sih = sih_box.text_frame
    tf_sih.word_wrap = True
    p_sih = tf_sih.paragraphs[0]
    p_sih.text = "SMART INDIA\nHACKATHON\n2026"
    p_sih.font.size = Pt(13)
    p_sih.font.bold = True
    p_sih.font.color.rgb = RGBColor(30, 58, 138)
    p_sih.font.name = "Arial"

    # -------------------------------------------------------------------------
    # LEFT COLUMN: CONCISE PROJECT SUMMARY BULLETS
    # -------------------------------------------------------------------------
    left_box = slide.shapes.add_textbox(Inches(0.4), Inches(1.2), Inches(4.3), Inches(5.8))
    tf_l = left_box.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = Inches(0)
    tf_l.margin_top = Inches(0)

    bullets = [
        "Maaraksha is an AI-powered maternal health early-warning and monitoring system designed to provide timely attention and support to expecting mothers, particularly in rural areas.",
        "In India, ASHA workers serve as vital grassroots links between pregnant women and the healthcare system. Maaraksha aims to digitalize and streamline this health monitoring process.",
        "Patients can create their health profile, upload medical reports, and check in daily through text, voice, or predefined options. AI analyzes this information to identify potential risks and provide guidance.",
        "Maaraksha digitalizes the records of expecting women assigned to ASHA workers, improving record maintenance, patient follow-up, and risk-based prioritization for timely assistance."
    ]

    for idx, b_text in enumerate(bullets):
        p_b = tf_l.paragraphs[0] if idx == 0 else tf_l.add_paragraph()
        p_b.text = "▪  " + b_text
        p_b.font.size = Pt(11.5)
        p_b.font.color.rgb = COLOR_TEXT
        p_b.font.name = "Arial"
        p_b.space_after = Pt(12)

    # -------------------------------------------------------------------------
    # CENTER & RIGHT COLUMN: SYSTEM WORKFLOW & DASHBOARD DIAGRAM
    # -------------------------------------------------------------------------
    if os.path.exists(diagram_img_path):
        # Positioned right next to text box
        slide.shapes.add_picture(diagram_img_path, Inches(4.8), Inches(1.15), Inches(8.133), Inches(5.1))

    # -------------------------------------------------------------------------
    # BOTTOM PARAGRAPH SUMMARY BELOW DIAGRAM
    # -------------------------------------------------------------------------
    bot_box = slide.shapes.add_textbox(Inches(4.8), Inches(6.35), Inches(8.133), Inches(0.95))
    tf_bot = bot_box.text_frame
    tf_bot.word_wrap = True
    tf_bot.margin_left = Inches(0)

    p_bot = tf_bot.paragraphs[0]
    p_bot.text = "The website uses role-based login for patients, ASHA workers, family members, PHC personnel, and district officers, allowing each to access relevant information. Higher authorities can monitor patient assignments, ASHA activities, and overall health-tracking progress."
    p_bot.font.size = Pt(11)
    p_bot.font.color.rgb = COLOR_TEXT
    p_bot.font.name = "Arial"

    # Save outputs
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    
    # Save copy to Maarakshak_Complete_System_Workflow.pptx as well
    copy_path = os.path.join(os.path.dirname(output_path), "Maarakshak_Complete_System_Workflow.pptx")
    prs.save(copy_path)
    
    print(f"SIH presentation generated at: {output_path}")

if __name__ == "__main__":
    out_ppt = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "Maaraksha_SIH_2026_Presentation.pptx"))
    img_in = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "public", "maaraksha_sih_flow_diagram.png"))
    create_sih_presentation(out_ppt, img_in)
